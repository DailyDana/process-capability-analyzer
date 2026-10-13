"""Tests for Gage R&R chart models, renderers and CSV/Excel export."""

from __future__ import annotations

import csv
import io
from pathlib import Path

import openpyxl
import pytest
from matplotlib.figure import Figure

from app.charts.gage_model import (
    COMPONENT_SOURCES,
    build_components_model,
    build_interaction_model,
    build_range_model,
)
from app.charts.gage_render import render_components, render_interaction, render_range_by_operator
from app.msa.models import GageMethod
from app.reporting.bundle import ExportMeta
from app.reporting.gage_report import gage_csv_string, gage_workbook_bytes
from app.services.gage_rr import columns_to_request, run_gage_study

EXAMPLE = Path(__file__).resolve().parent.parent / "examples" / "gage_rr_aiag.csv"


@pytest.fixture(scope="module", params=list(GageMethod), ids=lambda m: m.value)
def outcome(request):
    with open(EXAMPLE, newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    cols = [[r[k] for r in rows] for k in ("part", "operator", "value", "trial")]
    out = run_gage_study(columns_to_request(*cols, method=request.param, tolerance=8.0, unit="mm"))
    assert out.ok
    return out


class TestChartModels:
    def test_components_mirror_summary_rows(self, outcome):
        m = build_components_model(outcome.result)
        assert m.sources == COMPONENT_SOURCES
        assert m.pct_study_var[0] == outcome.result.pct_grr_study_var
        assert m.pct_tolerance is not None and m.pct_tolerance[0] == outcome.result.pct_grr_tolerance

    def test_range_model_flags_the_out_of_control_cell(self, outcome):
        m = build_range_model(outcome.study, outcome.result)
        b, p4 = m.operators.index("B"), m.parts.index("4")
        assert m.out_of_control[b][p4] is True
        assert sum(map(sum, m.out_of_control)) == 1
        assert m.ranges[b][p4] == pytest.approx(1.02)
        assert m.ucl == outcome.result.range_check.ucl

    def test_interaction_model_is_cell_means(self, outcome):
        m = build_interaction_model(outcome.study)
        assert m.averages[0][0] == pytest.approx((0.29 + 0.41 + 0.64) / 3)  # operator A, part 1

    def test_renderers_return_figures(self, outcome):
        import matplotlib.pyplot as plt

        figs = [
            render_components(build_components_model(outcome.result)),
            render_range_by_operator(build_range_model(outcome.study, outcome.result)),
            render_interaction(build_interaction_model(outcome.study)),
        ]
        assert all(isinstance(f, Figure) for f in figs)
        for f in figs:
            plt.close(f)


class TestCsv:
    def test_sections_and_precision(self, outcome):
        text = gage_csv_string(outcome.study, outcome.result, outcome.interpretation, ExportMeta(characteristic="Bore"))
        assert text == gage_csv_string(outcome.study, outcome.result, outcome.interpretation, ExportMeta(characteristic="Bore"))
        for header in ("# Metadata", "# Methodology", "# Gage R&R", "# Interpretation", "# Measurements"):
            assert header in text
        rows = list(csv.reader(io.StringIO(text)))
        grr = next(r for r in rows if r and r[0] == "Total Gage R&R")
        assert float(grr[2]) == outcome.result.row("Total Gage R&R").sd  # full precision round-trip
        meas_start = rows.index(["Part", "Operator", "Trial", "Value"])
        assert len(rows) - meas_start - 1 == 90

    def test_method_specific_sections(self, outcome):
        text = gage_csv_string(outcome.study, outcome.result)
        if outcome.result.method is GageMethod.ANOVA:
            assert "# ANOVA (full model, with interaction)" in text and "# ANOVA (interaction pooled)" in text
        else:
            assert "# Average & Range intermediates" in text and "# ANOVA" not in text
        assert "# Ranges above UCL" in text


class TestExcel:
    def test_sheets_and_values(self, outcome):
        wb = openpyxl.load_workbook(io.BytesIO(gage_workbook_bytes(outcome.study, outcome.result, outcome.interpretation)))
        expected = ["Gage R&R", "Interpretation", "Measurements"]
        if outcome.result.method is GageMethod.ANOVA:
            expected.insert(1, "ANOVA")
        assert wb.sheetnames == expected
        ws = wb["Gage R&R"]
        cells = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
        row = cells["Total Gage R&R"]
        assert ws.cell(row=row, column=3).value == pytest.approx(outcome.result.row("Total Gage R&R").sd)
        assert wb["Measurements"].max_row == 91
