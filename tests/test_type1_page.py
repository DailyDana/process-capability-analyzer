"""Type-1 page (AppTest), exports and PDF."""

from __future__ import annotations

import csv
import io
from pathlib import Path

import openpyxl
import pytest
from matplotlib.figure import Figure
from pypdf import PdfReader
from streamlit.testing.v1 import AppTest

from app.charts.type1_render import render_type1_run_chart
from app.msa.interpret import interpret_type1
from app.msa.type1 import Type1Request, type1_study
from app.reporting.bundle import ExportMeta
from app.reporting.msa_pdf import type1_pdf_bytes
from app.reporting.type1_report import type1_csv_string, type1_workbook_bytes

ROOT = Path(__file__).resolve().parent.parent
PAGE = str(ROOT / "app" / "ui" / "type1_page.py")


def _metric(at: AppTest, label: str) -> str:
    return next(m.value for m in at.metric if m.label == label)


class TestPage:
    def test_bosch_example_by_default(self):
        at = AppTest.from_file(PAGE, default_timeout=90)
        at.run()
        assert not at.exception
        assert [h.value for h in at.main.header] == [
            "1 · Readings and reference", "2 · Result", "3 · Run chart", "4 · Interpretation", "5 · Export"]
        assert _metric(at, "Cg") == "2.01" and _metric(at, "Cgk") == "1.64"
        assert any("Gauge CAPABLE" in s.value for s in at.success)

    def test_second_example_not_capable(self):
        at = AppTest.from_file(PAGE, default_timeout=90)
        at.run()
        at.radio(key="t1_source").set_value("Example: BPI Consulting (not capable)")
        at.run()
        assert not at.exception
        assert _metric(at, "Cg") == "0.34" and _metric(at, "Cgk") == "0.24"
        assert any("NOT CAPABLE" in e.value for e in at.error)

    def test_invalid_tolerance_reported(self):
        at = AppTest.from_file(PAGE, default_timeout=90)
        at.run()
        at.number_input(key=at.number_input[1].key).set_value(0.0)
        at.run()
        assert not at.exception
        assert any("Tolerance" in e.value for e in at.error)


@pytest.fixture(scope="module")
def bosch():
    with open(ROOT / "examples" / "type1_bosch.csv", newline="", encoding="utf-8") as fh:
        vals = tuple(float(r["value"]) for r in csv.DictReader(fh))
    r = type1_study(Type1Request(vals, 6.002, 0.060, resolution=0.001))
    return r, interpret_type1(r)


class TestExports:
    def test_csv(self, bosch):
        r, rep = bosch
        text = type1_csv_string(r, rep, ExportMeta(characteristic="W450"))
        rows = list(csv.reader(io.StringIO(text)))
        cg = next(row for row in rows if row and row[0] == "Cg")
        assert float(cg[1]) == r.cg
        assert rows.index(["Reading", "Value"]) == len(rows) - 51

    def test_excel(self, bosch):
        r, rep = bosch
        wb = openpyxl.load_workbook(io.BytesIO(type1_workbook_bytes(r, rep)))
        assert wb.sheetnames == ["Type-1 study", "Interpretation", "Readings"]
        assert wb["Readings"].max_row == 51

    def test_pdf(self, bosch):
        r, rep = bosch
        pdf = type1_pdf_bytes(r, rep, ExportMeta(characteristic="W450 diameter", unit="mm"))
        text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
        for needle in ("Type-1 Gauge Study Report", "W450 diameter", "2.01", "1.64", "type1_bias_significant",
                       "rec_type1_proceed", "5. Readings", "5.999"):
            assert needle in text
        assert "&amp;" not in text and "&lt;" not in text

    def test_run_chart(self, bosch):
        import matplotlib.pyplot as plt

        fig = render_type1_run_chart(bosch[0], unit="mm")
        assert isinstance(fig, Figure)
        plt.close(fig)
