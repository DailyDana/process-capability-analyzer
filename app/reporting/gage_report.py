"""CSV and Excel export of a Gage R&R analysis.

Serialization only: every number comes from the :class:`GageRRResult` /
:class:`GageStudy` already computed; nothing is recalculated here. Numbers are
written at full precision (CSV via ``repr``), matching the capability exports.
"""

from __future__ import annotations

import csv
import io
from typing import Optional

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.interpret.models import InterpretationReport
from app.msa.anova import AnovaResult, AnovaTable
from app.msa.models import GageMethod, GageRRResult, GageStudy
from app.msa.xbar_r import AverageRangeDetail
from app.reporting.bundle import ExportMeta
from app.reporting.csv_report import _num as _csv_num
from app.reporting.csv_report import _s as _csv_s
from app.reporting.excel_report import _SEV_FILL, _autosize, _header_row
from app.reporting.excel_report import _num as _xl_num

__all__ = ["gage_csv_string", "gage_workbook_bytes", "methodology_rows"]

_METHOD_LABEL = {GageMethod.ANOVA: "ANOVA", GageMethod.AVERAGE_RANGE: "Average & Range (AIAG)"}


def methodology_rows(result: GageRRResult) -> list[tuple[str, object]]:
    rows: list[tuple[str, object]] = [
        ("Method", _METHOD_LABEL[result.method]),
        ("Parts", result.n_parts),
        ("Operators", result.n_operators),
        ("Trials", result.n_trials),
        ("Study variation multiplier (k)", result.k),
        ("Tolerance (USL - LSL)", result.tolerance),
    ]
    if isinstance(result.anova, AnovaResult):
        a = result.anova
        rows.append(("Interaction pooling alpha", a.alpha_pool))
        rows.append(("Interaction p-value", a.interaction_p))
        rows.append(("Interaction pooled", a.interaction_pooled))
    rows.append(("ndc", result.ndc))
    rows.append(("Range chart R-bar", result.range_check.r_bar))
    rows.append(("Range chart UCL", result.range_check.ucl))
    rows.append(("Ranges above UCL", len(result.range_check.out_of_control)))
    return rows


def _anova_rows(table: AnovaTable) -> list[list[object]]:
    return [[r.source, r.df, r.ss, r.ms, r.f, r.p] for r in table.rows]


_SUMMARY_HEADERS = ["Source", "Variance", "StdDev", "Study Var (k*SD)", "% Contribution", "% Study Var", "% Tolerance"]
_ANOVA_HEADERS = ["Source", "DF", "SS", "MS", "F", "P"]


def _summary_rows(result: GageRRResult) -> list[list[object]]:
    return [
        [r.source, r.variance, r.sd, r.study_var, r.pct_contribution, r.pct_study_var, r.pct_tolerance]
        for r in result.rows
    ]


def _xbar_r_rows(d: AverageRangeDetail) -> list[tuple[str, object]]:
    return [
        ("R-bar-bar", d.r_bar_bar),
        ("X-bar diff (max - min operator average)", d.x_bar_diff),
        ("R_p (max - min part average)", d.r_p),
        ("K1", d.k1),
        ("K2", d.k2),
        ("K3", d.k3),
        ("AV radicand", d.av_radicand),
    ]


def _fmt_csv(v: object) -> str:
    if isinstance(v, bool) or v is None or isinstance(v, str):
        return _csv_s(v)
    if isinstance(v, (int, float)):
        return _csv_num(float(v)) if isinstance(v, float) else str(v)
    return _csv_s(v)


def gage_csv_string(
    study: GageStudy,
    result: GageRRResult,
    interpretation: Optional[InterpretationReport] = None,
    meta: ExportMeta = ExportMeta(),
    warnings: tuple[str, ...] = (),
) -> str:
    """Deterministic CSV: metadata, methodology, table(s), interpretation, data."""
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")

    def section(title: str) -> None:
        if buf.tell():
            w.writerow([])
        w.writerow([f"# {title}"])

    section("Metadata")
    w.writerow(["Field", "Value"])
    for label, value in meta.as_rows():
        w.writerow([label, _csv_s(value)])

    section("Methodology")
    w.writerow(["Field", "Value"])
    for label, value in methodology_rows(result):
        w.writerow([label, _fmt_csv(value)])

    section("Gage R&R")
    w.writerow(_SUMMARY_HEADERS)
    for row in _summary_rows(result):
        w.writerow([row[0], *(_fmt_csv(v) for v in row[1:])])

    if isinstance(result.anova, AnovaResult):
        section("ANOVA (full model, with interaction)")
        w.writerow(_ANOVA_HEADERS)
        for row in _anova_rows(result.anova.full):
            w.writerow([_fmt_csv(v) for v in row])
        if result.anova.reduced is not None:
            section("ANOVA (interaction pooled)")
            w.writerow(_ANOVA_HEADERS)
            for row in _anova_rows(result.anova.reduced):
                w.writerow([_fmt_csv(v) for v in row])
    if isinstance(result.average_range, AverageRangeDetail):
        section("Average & Range intermediates")
        w.writerow(["Field", "Value"])
        for label, value in _xbar_r_rows(result.average_range):
            w.writerow([label, _fmt_csv(value)])

    if result.range_check.out_of_control:
        section("Ranges above UCL")
        w.writerow(["Part", "Operator", "Range"])
        for p, o, r in result.range_check.out_of_control:
            w.writerow([p, o, _fmt_csv(r)])

    if warnings:
        section("Data warnings")
        for msg in warnings:
            w.writerow([msg])

    if interpretation is not None:
        section("Interpretation")
        w.writerow(["Headline", interpretation.headline])
        w.writerow(["Kind", "Code", "Severity", "Text"])
        for c in interpretation.conclusions:
            w.writerow(["conclusion", c.code, c.severity.value, c.text])
        for r in interpretation.recommendations:
            w.writerow(["recommendation", r.code, r.severity.value, r.text])

    section("Measurements")
    w.writerow(["Part", "Operator", "Trial", "Value"])
    for i, part in enumerate(study.part_labels):
        for j, op in enumerate(study.operator_labels):
            for k in range(study.n_trials):
                w.writerow([part, op, k + 1, _fmt_csv(float(study.values[i, j, k]))])
    return buf.getvalue()


def _xl(v: object) -> object:
    """Excel cell value: floats via the shared formatter, None as "n/a"."""
    if v is None or (isinstance(v, float)):
        return _xl_num(v)  # type: ignore[arg-type]
    return v


def _write_kv(ws: Worksheet, start: int, rows: list[tuple[str, object]]) -> int:
    _header_row(ws, start, ["Field", "Value"])
    r = start + 1
    for label, value in rows:
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=_xl(value))
        r += 1
    return r + 1


def _write_table(ws: Worksheet, start: int, headers: list[str], rows: list[list[object]]) -> int:
    _header_row(ws, start, headers)
    r = start + 1
    for row in rows:
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=_xl(v))
        r += 1
    return r + 1


def gage_workbook_bytes(
    study: GageStudy,
    result: GageRRResult,
    interpretation: Optional[InterpretationReport] = None,
    meta: ExportMeta = ExportMeta(),
    warnings: tuple[str, ...] = (),
) -> bytes:
    """Excel workbook: Summary, ANOVA (when used), Interpretation, Measurements."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Gage R&R"
    r = 1
    if meta.as_rows():
        r = _write_kv(ws, r, meta.as_rows())
    r = _write_kv(ws, r, methodology_rows(result))
    r = _write_table(ws, r, _SUMMARY_HEADERS, _summary_rows(result))
    if isinstance(result.average_range, AverageRangeDetail):
        r = _write_kv(ws, r, _xbar_r_rows(result.average_range))
    for msg in warnings:
        ws.cell(row=r, column=1, value=f"Data warning: {msg}")
        r += 1
    _autosize(ws)

    if isinstance(result.anova, AnovaResult):
        wa = wb.create_sheet("ANOVA")
        ra = 1
        wa.cell(row=ra, column=1, value="Full model (with interaction)")
        ra = _write_table(wa, ra + 1, _ANOVA_HEADERS, _anova_rows(result.anova.full))
        if result.anova.reduced is not None:
            wa.cell(row=ra, column=1, value="Interaction pooled into repeatability")
            _write_table(wa, ra + 1, _ANOVA_HEADERS, _anova_rows(result.anova.reduced))
        _autosize(wa)

    if interpretation is not None:
        wi = wb.create_sheet("Interpretation")
        wi.cell(row=1, column=1, value=interpretation.headline)
        _header_row(wi, 3, ["Kind", "Code", "Severity", "Text"])
        ri = 4
        items = [("conclusion", c.code, c.severity, c.text) for c in interpretation.conclusions]
        items += [("recommendation", x.code, x.severity, x.text) for x in interpretation.recommendations]
        for kind, code, sev, text in items:
            for c, v in enumerate((kind, code, sev.value, text), start=1):
                cell = wi.cell(row=ri, column=c, value=v)
                if c == 3 and sev.value in _SEV_FILL:
                    cell.fill = _SEV_FILL[sev.value]
            ri += 1
        _autosize(wi, max_width=100)

    wm = wb.create_sheet("Measurements")
    _header_row(wm, 1, ["Part", "Operator", "Trial", "Value"])
    rm = 2
    for i, part in enumerate(study.part_labels):
        for j, op in enumerate(study.operator_labels):
            for k in range(study.n_trials):
                for c, v in enumerate((part, op, k + 1, float(study.values[i, j, k])), start=1):
                    wm.cell(row=rm, column=c, value=v)
                rm += 1
    _autosize(wm)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
