"""CSV and Excel export of a Type-1 gauge study (serialization only)."""

from __future__ import annotations

import csv
import io
from typing import Optional

from openpyxl import Workbook

from app.interpret.models import InterpretationReport
from app.msa.type1 import Type1Result
from app.reporting.bundle import ExportMeta
from app.reporting.csv_report import _num as _csv_num
from app.reporting.csv_report import _s as _csv_s
from app.reporting.excel_report import _SEV_FILL, _autosize, _header_row
from app.reporting.excel_report import _num as _xl_num

__all__ = ["result_rows", "type1_csv_string", "type1_workbook_bytes"]


def result_rows(r: Type1Result) -> list[tuple[str, object]]:
    """(label, value) rows shared by CSV, Excel, PDF and the UI."""
    tmin = r.min_tolerances()
    rows: list[tuple[str, object]] = [
        ("Readings (n)", r.n),
        ("Reference value x_m", r.reference),
        ("Tolerance T", r.tolerance),
        ("K (% of tolerance)", r.k_percent),
        ("L (sigma multiplier)", r.l_sigma),
        ("Resolution", r.resolution),
        ("Resolution, % of T", r.pct_resolution),
        ("Mean", r.mean),
        ("Standard deviation s", r.sd),
        ("Bias (mean - x_m)", r.bias),
        ("Cg", r.cg),
        ("Cg confidence interval", None if r.cg_ci is None else f"{r.cg_ci[0]:.4g} - {r.cg_ci[1]:.4g}"),
        ("Cgk", r.cgk),
        ("Cgk confidence interval", None if r.cgk_ci is None else f"{r.cgk_ci[0]:.4g} - {r.cgk_ci[1]:.4g}"),
        ("Confidence level", 1 - r.alpha),
        ("%Var (repeatability)", r.pct_var_rep),
        ("%Var (repeatability and bias)", r.pct_var_rep_bias),
        ("Bias t statistic", r.t_stat),
        ("Bias p-value", r.p_value),
        ("Bias significant", r.bias_significant),
        ("Bias criterion |bias|/s", r.bias_criterion),
        ("Smallest tolerance for Cg >= 1.33", tmin["Cg"]),
        ("Smallest tolerance for Cgk >= 1.33", tmin["Cgk"]),
        ("Smallest tolerance for resolution <= 5%", tmin["RE"]),
    ]
    return rows


def _csv_val(v: object) -> str:
    if isinstance(v, float):
        return _csv_num(v)
    return _csv_s(v)


def type1_csv_string(r: Type1Result, interpretation: Optional[InterpretationReport] = None,
                     meta: ExportMeta = ExportMeta(), warnings: tuple[str, ...] = ()) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["# Metadata"])
    w.writerow(["Field", "Value"])
    for k, v in meta.as_rows():
        w.writerow([k, _csv_s(v)])
    w.writerow([])
    w.writerow(["# Type-1 gauge study"])
    w.writerow(["Field", "Value"])
    for k, v in result_rows(r):
        w.writerow([k, _csv_val(v)])
    if warnings:
        w.writerow([])
        w.writerow(["# Data warnings"])
        for m in warnings:
            w.writerow([m])
    if interpretation is not None:
        w.writerow([])
        w.writerow(["# Interpretation"])
        w.writerow(["Headline", interpretation.headline])
        w.writerow(["Kind", "Code", "Severity", "Text"])
        for c in interpretation.conclusions:
            w.writerow(["conclusion", c.code, c.severity.value, c.text])
        for x in interpretation.recommendations:
            w.writerow(["recommendation", x.code, x.severity.value, x.text])
    w.writerow([])
    w.writerow(["# Readings"])
    w.writerow(["Reading", "Value"])
    for i, v in enumerate(r.values, start=1):
        w.writerow([i, _csv_num(v)])
    return buf.getvalue()


def type1_workbook_bytes(r: Type1Result, interpretation: Optional[InterpretationReport] = None,
                         meta: ExportMeta = ExportMeta(), warnings: tuple[str, ...] = ()) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Type-1 study"
    _header_row(ws, 1, ["Field", "Value"])
    row = 2
    for k, v in meta.as_rows() + result_rows(r):
        ws.cell(row=row, column=1, value=k)
        ws.cell(row=row, column=2, value=_xl_num(v) if isinstance(v, float) or v is None else v)
        row += 1
    for m in warnings:
        ws.cell(row=row + 1, column=1, value=f"Data warning: {m}")
        row += 1
    _autosize(ws)
    if interpretation is not None:
        wi = wb.create_sheet("Interpretation")
        wi.cell(row=1, column=1, value=interpretation.headline)
        _header_row(wi, 3, ["Kind", "Code", "Severity", "Text"])
        ri = 4
        items = [("conclusion", c.code, c.severity, c.text) for c in interpretation.conclusions]
        items += [("recommendation", x.code, x.severity, x.text) for x in interpretation.recommendations]
        for kind, code, sev, text in items:
            for col, val in enumerate((kind, code, sev.value, text), start=1):
                cell = wi.cell(row=ri, column=col, value=val)
                if col == 3 and sev.value in _SEV_FILL:
                    cell.fill = _SEV_FILL[sev.value]
            ri += 1
        _autosize(wi, max_width=100)
    wr = wb.create_sheet("Readings")
    _header_row(wr, 1, ["Reading", "Value"])
    for i, v in enumerate(r.values, start=1):
        wr.cell(row=i + 1, column=1, value=i)
        wr.cell(row=i + 1, column=2, value=v)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
